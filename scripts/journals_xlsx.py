"""Write docs/IEEE_Journals_APC.xlsx from data/journals.yml (called by build.py).

APCs are live formulas that read the 'Pricing rules' sheet, so a reader can change a price
there and the whole workbook updates. The file is written deterministically (fixed zip
timestamps and document dates) so the daily rebuild only commits it when the data changes.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
INPUT = Font(name=ARIAL, size=10, color="0000FF", bold=True)
LINK = Font(name=ARIAL, size=10, color="0563C1", underline="single")
TITLE = Font(name=ARIAL, size=14, bold=True)
NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
WARN = Font(name=ARIAL, size=10, color="C00000")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
UNCONFIRMED = "Open-access model not confirmed for 2026 — check the journal's author page."


def _header(ws, row: int, labels: list[str]) -> None:
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(wrap_text=True, vertical="center")


def _table(ws, headers, rows, widths, name):
    _header(ws, 1, headers)
    for r, vals in enumerate(rows, 2):
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.alignment, c.border = BODY, WRAP, THIN
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    t = Table(displayName=name, ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def build_xlsx(pricing: dict, journals: list[dict], path: Path) -> None:
    year = pricing["year"]
    models = pricing["models"]
    disc = pricing["discounts"]
    wb = Workbook()

    # Read me
    rm = wb.active
    rm.title = "Read me"
    lines = [
        (f"IEEE journals — open-access model, APC {year} and time to first decision", TITLE),
        ("• Hybrid = subscription journal. Publishing is FREE unless you choose open access (then you pay the APC). Overlength page charges may still apply.", BODY),
        ("• Fully open access = every accepted article pays the APC. There is no free option (waivers exist for low-income countries).", BODY),
        ("• APCs are formulas that read the 'Pricing rules' sheet — change a blue cell there and every journal updates. A few titles have their own price (set in the data).", BODY),
        ("• Decision times: IEEE does not publish a time to first decision for most journals. Figures appear only where an IEEE society or a named source gives one; the source type is shown next to it.", BODY),
        ("• Red notes: open-access model not confirmed in 2026 sources — check the journal's author page.", BODY),
        ("Generated from data/journals.yml in github.com/aionndc22/ieee-conference-deadlines — contributions welcome. Not affiliated with IEEE.", NOTE),
    ]
    for r, (txt, f) in enumerate(lines, 1):
        c = rm.cell(row=r, column=1, value=txt)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")
    rm.column_dimensions["A"].width = 120

    # Pricing rules (inputs)
    pr = wb.create_sheet("Pricing rules")
    pr["A1"] = f"IEEE open-access pricing — {year} submissions (last checked {pricing.get('verified', '')})"
    pr["A1"].font = TITLE
    pr["A2"] = "Blue cells are inputs; the Journals sheet recalculates from them."
    pr["A2"].font = NOTE
    _header(pr, 3, ["Publishing model", f"APC {year} (US$)", "If you don't choose open access", "What it means"])
    model_names = list(models)
    for r, m in enumerate(model_names, 4):
        spec = models[m]
        pr.cell(row=r, column=1, value=m).font = BOLD
        c = pr.cell(row=r, column=2, value=spec["apc"])
        c.font, c.number_format, c.fill = INPUT, "$#,##0", YELLOW
        pr.cell(row=r, column=3, value=spec.get("without_oa", "")).font = BODY
        pr.cell(row=r, column=4, value=spec.get("meaning", "")).font = BODY
    last_model = 3 + len(model_names)
    d0 = last_model + 2
    pr.cell(row=d0, column=1, value="Discounts").font = BOLD
    pr.cell(row=d0 + 1, column=1, value="IEEE member discount").font = BODY
    pr.cell(row=d0 + 2, column=1, value="IEEE Society member discount").font = BODY
    for r, key in ((d0 + 1, "ieee_member"), (d0 + 2, "society_member")):
        c = pr.cell(row=r, column=2, value=disc[key])
        c.font, c.number_format, c.fill = INPUT, "0%", YELLOW
    pr.cell(row=d0 + 2, column=3, value="Discounts do not stack; students excluded.").font = BODY
    s0 = d0 + 4
    pr.cell(row=s0, column=1, value="Sources").font = BOLD
    for i, src in enumerate(pricing.get("sources", []), s0 + 1):
        c = pr.cell(row=i, column=1, value=src)
        c.font, c.hyperlink = LINK, src
    for row in pr.iter_rows(min_row=4, max_row=d0 + 2):
        for c in row:
            c.alignment = WRAP
    for col, w in zip("ABCD", (34, 16, 40, 80)):
        pr.column_dimensions[col].width = w
    rule_models = f"'Pricing rules'!$A$4:$A${last_model}"
    rule_apc = f"'Pricing rules'!$B$4:$B${last_model}"
    soc = f"'Pricing rules'!$B${d0 + 2}"

    def apc_formula(j: dict, r: int) -> str | int:
        if "apc" in j:  # title with its own price
            return int(j["apc"])
        return f"=INDEX({rule_apc},MATCH(E{r},{rule_models},0))"

    def rows_for(items: list[dict]) -> list[list]:
        out = []
        for r, j in enumerate(items, 2):
            note = " ".join(x for x in (UNCONFIRMED if j.get("model_unconfirmed") else "", j.get("notes", "")) if x)
            out.append([
                j["track"], j["name"], j["acronym"], j["type"], j["oa_model"],
                apc_formula(j, r), f"=ROUND(F{r}*(1-{soc}),0)",
                models[j["oa_model"]].get("without_oa", ""),
                j.get("decision_time", "Not published"), j.get("decision_source", ""),
                j.get("decision_url", ""), "AESS" if j.get("aess") else "", note,
            ])
        return out

    headers = ["Track", "Journal", "Acronym", "Type", "OA model", f"APC {year} (US$)",
               "APC for society member (US$)", "Cost without OA", "Avg. time to first decision",
               "Decision-time source", "Source link", "Society", "Notes"]
    widths = (30, 52, 11, 13, 17, 13, 15, 24, 30, 26, 34, 9, 48)
    regular = [j for j in journals if j["type"] != "Magazine"]
    mags = [j for j in journals if j["type"] == "Magazine"]
    for title, items, tname in (("Journals", regular, "Journals"), ("Magazines", mags, "Magazines")):
        ws = wb.create_sheet(title)
        _table(ws, headers, rows_for(items), widths, tname)
        for r, j in enumerate(items, 2):
            ws.cell(row=r, column=2).font = BOLD
            for col in (6, 7):
                ws.cell(row=r, column=col).number_format = "$#,##0"
            link = ws.cell(row=r, column=11)
            if link.value:
                link.hyperlink, link.font = link.value, LINK
            if j.get("model_unconfirmed"):
                ws.cell(row=r, column=13).font = WARN

    # Summary by track (formulas over the Journals sheet)
    sm = wb.create_sheet("Summary by track", 2)
    sm["A1"] = "IEEE journals by track — counts by publishing model"
    sm["A1"].font = TITLE
    _header(sm, 3, ["Track", "Journals", "Hybrid", "Fully open access", "With a decision-time figure"])
    n = len(regular) + 1
    tracks = list(dict.fromkeys(j["track"] for j in regular))
    for r, tr in enumerate(tracks, 4):
        sm.cell(row=r, column=1, value=tr).font = BODY
        sm.cell(row=r, column=2, value=f"=COUNTIF(Journals!$A$2:$A${n},A{r})").font = BODY
        sm.cell(row=r, column=3, value=f'=COUNTIFS(Journals!$A$2:$A${n},A{r},Journals!$E$2:$E${n},"Hybrid")').font = BODY
        sm.cell(row=r, column=4, value=f'=COUNTIFS(Journals!$A$2:$A${n},A{r},Journals!$E$2:$E${n},"Fully open access")').font = BODY
        sm.cell(row=r, column=5, value=f'=COUNTIFS(Journals!$A$2:$A${n},A{r},Journals!$I$2:$I${n},"<>Not published")').font = BODY
    tot = 4 + len(tracks)
    sm.cell(row=tot, column=1, value="Total").font = BOLD
    for col in range(2, 6):
        L = get_column_letter(col)
        sm.cell(row=tot, column=col, value=f"=SUM({L}4:{L}{tot - 1})").font = BOLD
    for col, w in zip("ABCDE", (40, 12, 12, 18, 26)):
        sm.column_dimensions[col].width = w

    # Deterministic output: fixed document dates and zip timestamps.
    fixed = dt.datetime(2026, 1, 1)
    wb.properties.created = wb.properties.modified = fixed
    wb.properties.creator = "ieee-conference-deadlines"
    buf = io.BytesIO()
    wb.save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, \
            zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps the save time here
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, data)
